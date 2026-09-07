"""Document converters. COPY-THEN-DIVERGE from jester-2.1/ingest/conv_pptx.py
and conv_img.py per DR-033(c) (fork point: jester-2.1 as of 2026-09-07,
thread D1 ingest) -- reimplemented here, not imported, not extracted to a
shared package (PORTFOLIO.md §5).

Divergence from 2.x: 2.x's ConversionResult/version_meta/manifest apparatus
is audit-evidence machinery (approval dates, confidence scoring for the
assessor's evidence_strength field) that 1.x's retrieval-and-trigger use case
does not need. This module returns a plain (text, status, reason) tuple --
deliberately smaller than 2.x's, per the no-premature-abstraction default.
Only pptx and image conversion are implemented: DR-031's survey of the
mounted Jester_IN volume (thread D1) found no pdf/docx/xlsx/msg/html/eml
files on it, so those converters were not copied across. Extending this
module when such a file class actually appears is expected, not designed
speculatively now.

Licence note (DR-006 D4 / CLAUDE.md, reverified for this session rather than
assumed): python-pptx (MIT), Pillow (MIT-CMU), pytesseract (Apache-2.0). No
PyMuPDF, no extract-msg anywhere in this module -- neither format is present
in the surveyed corpus so neither dependency was ever pulled in.
"""
from __future__ import annotations

from pathlib import Path

from PIL import ExifTags, Image
from pptx import Presentation

from . import ocr

OK = "ok"
DEGRADED = "degraded"
UNPROCESSED = "unprocessed"


def _table_to_markdown(table) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""
    lines = ["| " + " | ".join(rows[0]) + " |",
             "| " + " | ".join("---" for _ in rows[0]) + " |"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def _slide_title(slide) -> str:
    title_shape = slide.shapes.title
    if title_shape is not None and title_shape.has_text_frame:
        return title_shape.text_frame.text.strip()
    return ""


def _slide_body_parts(slide, title: str) -> list[str]:
    parts: list[str] = []
    for shape in slide.shapes:
        if shape.has_table:
            md = _table_to_markdown(shape.table)
            if md:
                parts.append(md)
        elif shape.has_text_frame:
            text = shape.text_frame.text.strip()
            if text and text != title:
                parts.append(text)
    return parts


def _slide_notes(slide) -> str:
    if slide.has_notes_slide:
        return slide.notes_slide.notes_text_frame.text.strip()
    return ""


def convert_pptx(path: Path) -> tuple[str, str, str]:
    """Returns (text, status, reason)."""
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        return "", UNPROCESSED, f"conversion-error: {type(exc).__name__}: {exc}"

    sections: list[str] = []
    for i, slide in enumerate(presentation.slides, start=1):
        title = _slide_title(slide)
        heading = f"## Slide {i}: {title}" if title else f"## Slide {i}"
        section = [heading, *_slide_body_parts(slide, title)]
        notes = _slide_notes(slide)
        if notes:
            section.append(f"[notes] {notes}")
        sections.append("\n\n".join(section))

    text = "\n\n".join(sections)
    if text.strip():
        return text, OK, ""
    return "", DEGRADED, "no-extractable-text"


_EXIF_TAG_IDS = {name: tag_id for tag_id, name in ExifTags.TAGS.items()}


def convert_image(path: Path) -> tuple[str, str, str]:
    """Returns (text, status, reason). OCR via ingest.ocr; see that module's
    docstring for the tesseract-missing degrade path."""
    try:
        image = Image.open(path)
        image.load()
    except Exception as exc:
        return "", UNPROCESSED, f"conversion-error: {type(exc).__name__}: {exc}"

    if ocr.tesseract_missing():
        return "", DEGRADED, "ocr-unavailable: tesseract binary not found"

    text, _confidence = ocr.run_ocr(image)
    if text.strip():
        return text, OK, ""
    return "", DEGRADED, "no-text-detected"
